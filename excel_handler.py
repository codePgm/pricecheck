"""
엑셀 파일 읽기/쓰기 처리 모듈
"""

import openpyxl
from openpyxl.styles import PatternFill
from pathlib import Path
import shutil
from datetime import datetime
import config

class ExcelHandler:
    def __init__(self, excel_path):
        self.excel_path = Path(excel_path)
        self.workbook = None
        self.worksheet = None
        self.data = {}
        self.headers = None
        
    def load(self):
        """엑셀 파일 로드"""
        if not self.excel_path.exists():
            raise FileNotFoundError(f"엑셀 파일을 찾을 수 없습니다: {self.excel_path}")
        
        self.workbook = openpyxl.load_workbook(self.excel_path)
        self.worksheet = self.workbook.active
        
        # 데이터 읽기
        for idx, row in enumerate(self.worksheet.iter_rows(values_only=True), 1):
            if idx == 1:
                self.headers = row
                continue
            
            if row[0]:  # 아이템 이름이 있는 경우만
                item_name = row[0]
                # 가격은 마지막 컬럼 (숫자 형태)
                price = row[-1] if len(row) > 3 else None
                
                self.data[item_name] = {
                    'row_num': idx,
                    'category': row[1] if len(row) > 1 else None,
                    'display_price': row[2] if len(row) > 2 else None,
                    'price': price,
                    'original_price': price  # 원본 가격 보관
                }
        
        return self.data
    
    def get_price(self, item_name):
        """특정 아이템의 가격 조회"""
        if item_name in self.data:
            return self.data[item_name]['price']
        return None
    
    def update_price(self, item_name, new_price, highlight=True):
        """가격 업데이트"""
        if item_name not in self.data:
            return False
        
        row_num = self.data[item_name]['row_num']
        
        # 가격 컬럼은 마지막 컬럼 (D열, 인덱스 4)
        price_col = 4
        
        # 엑셀에 값 업데이트
        self.worksheet.cell(row=row_num, column=price_col, value=new_price)
        
        # 표시 가격도 업데이트 (C열, 인덱스 3)
        display_price = self._format_price_display(new_price)
        self.worksheet.cell(row=row_num, column=3, value=display_price)
        
        # 하이라이트
        if highlight:
            yellow_fill = PatternFill(start_color=config.HIGHLIGHT_COLOR,
                                     end_color=config.HIGHLIGHT_COLOR,
                                     fill_type="solid")
            self.worksheet.cell(row=row_num, column=price_col).fill = yellow_fill
            self.worksheet.cell(row=row_num, column=3).fill = yellow_fill
        
        # 메모리 데이터도 업데이트
        self.data[item_name]['price'] = new_price
        self.data[item_name]['display_price'] = display_price
        
        return True
    
    def _format_price_display(self, price):
        """가격을 읽기 쉬운 형태로 변환"""
        if price >= 100000000:  # 1억 이상
            eok = price // 100000000
            man = (price % 100000000) // 10000
            if man > 0:
                return f"{eok}억 {man}만"
            else:
                return f"{eok}억"
        elif price >= 10000:  # 1만 이상
            man = price // 10000
            rest = price % 10000
            if rest > 0:
                return f"{man}만 {rest}"
            else:
                return f"{man}만"
        else:
            return str(price)
    
    def create_backup(self):
        """백업 파일 생성"""
        if not config.CREATE_BACKUP:
            return None
        
        backup_name = config.BACKUP_EXCEL_PATTERN.format(
            original_name=self.excel_path.stem
        )
        backup_path = self.excel_path.parent / backup_name
        
        shutil.copy2(self.excel_path, backup_path)
        return backup_path
    
    def save(self, output_path=None):
        """엑셀 파일 저장"""
        if output_path is None:
            output_name = config.OUTPUT_EXCEL_PATTERN.format(
                original_name=self.excel_path.stem
            )
            # 원본 엑셀 파일과 같은 폴더에 저장
            output_path = self.excel_path.parent / output_name
        
        self.workbook.save(output_path)
        return output_path
    
    def get_all_items(self):
        """모든 아이템 목록 반환"""
        return list(self.data.keys())
    
    def get_items_by_series(self, series_keyword):
        """특정 시리즈의 아이템들만 반환"""
        items = []
        for item_name in self.data.keys():
            if series_keyword in item_name:
                items.append(item_name)
        return items
    
    def close(self):
        """워크북 닫기"""
        if self.workbook:
            self.workbook.close()
